from __future__ import annotations

import io
import json
from copy import deepcopy
from datetime import datetime
from typing import Any

from fpdf import FPDF
from fpdf.enums import XPos, YPos


def shape_mcq_for_mode(mcq_data: dict[str, Any], mode: str) -> dict[str, Any]:
    """Return export-safe MCQ payload by mode.

    raw mode: include all fields.
    test mode: strip right_answer and explanation from every question.
    """
    shaped = deepcopy(mcq_data)

    if mode != "test":
        return shaped

    for question in shaped.get("questions", []):
        question.pop("right_answer", None)
        question.pop("explanation", None)

    return shaped


def format_mcq_as_pdf(mcq_data: dict[str, Any], mode: str) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # TODO: Use a custom header/footer FPDF class once branding requirements are finalized.
    pdf.set_font("Helvetica", "B", 14)
    pdf.multi_cell(
        0,
        10,
        f"MCQ {datetime.now():%d/%m/%y}",
        new_x=XPos.LMARGIN,
        new_y=YPos.NEXT,
    )
    pdf.ln(2)

    questions = mcq_data.get("questions", [])

    for index, question in enumerate(questions, start=1):
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(
            0,
            8,
            f"{index}. {question.get('question', '')}",
            new_x=XPos.LMARGIN,
            new_y=YPos.NEXT,
        )

        pdf.set_font("Helvetica", "", 11)
        for option in question.get("options", []):
            option_key = option.get("key", "")
            option_text = option.get("text", "")
            pdf.multi_cell(
                0,
                7,
                f"- {option_key}. {option_text}",
                new_x=XPos.LMARGIN,
                new_y=YPos.NEXT,
            )

        if mode == "raw":
            answer = question.get("right_answer", "")
            explanation = question.get("explanation", "")
            pdf.set_font("Helvetica", "I", 11)
            pdf.multi_cell(
                0,
                7,
                f"Correct answer: {answer}",
                new_x=XPos.LMARGIN,
                new_y=YPos.NEXT,
            )
            pdf.multi_cell(
                0,
                7,
                f"Explanation: {explanation}",
                new_x=XPos.LMARGIN,
                new_y=YPos.NEXT,
            )

        pdf.ln(2)

    output = pdf.output(dest="S")
    return output.encode("latin-1") if isinstance(output, str) else bytes(output)


def format_mcq_as_docx(mcq_data: dict[str, Any], mode: str) -> bytes:
    try:
        from docx import Document
        from docx.shared import RGBColor
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "DOCX export requires 'python-docx'. Install dependencies and rebuild the API image."
        ) from exc

    document = Document()
    document.add_heading(f"MCQ {datetime.now():%d/%m/%y}", level=1)

    questions = mcq_data.get("questions", [])
    for index, question in enumerate(questions, start=1):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"{index}. {question.get('question', '')}")
        run.bold = True

        for option in question.get("options", []):
            option_para = document.add_paragraph(
                f"{option.get('key', '')}. {option.get('text', '')}"
            )
            option_para.paragraph_format.left_indent = 457200

        if mode == "raw":
            answer_run = document.add_paragraph().add_run(
                f"Correct answer: {question.get('right_answer', '')}"
            )
            answer_run.italic = True
            answer_run.font.color.rgb = RGBColor(0x00, 0x80, 0x00)

            explanation_run = document.add_paragraph().add_run(
                f"Explanation: {question.get('explanation', '')}"
            )
            explanation_run.italic = True
            explanation_run.font.color.rgb = RGBColor(0x00, 0x80, 0x00)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def format_mcq_as_markdown(mcq_data: dict[str, Any], mode: str) -> bytes:
    lines: list[str] = [f"# MCQ {datetime.now():%d/%m/%y}", ""]

    questions = mcq_data.get("questions", [])
    for index, question in enumerate(questions, start=1):
        lines.append(f"{index}. {question.get('question', '')}")
        for option in question.get("options", []):
            lines.append(f"   - {option.get('key', '')}. {option.get('text', '')}")

        if mode == "raw":
            lines.append(f"   - Correct answer: {question.get('right_answer', '')}")
            lines.append(f"   - Explanation: {question.get('explanation', '')}")

        lines.append("")

    return "\n".join(lines).encode("utf-8")


def format_mcq_as_json(mcq_data: dict[str, Any], mode: str) -> bytes:
    questions = mcq_data.get("questions", [])
    payload = {
        "title": mcq_data.get("test_name", "MCQ Export"),
        "total_count": len(questions),
        "mode": mode,
        "questions": questions,
    }
    return json.dumps(payload, indent=2, default=str).encode("utf-8")


def build_export_filename(mode: str, export_format: str) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"mcq_{mode}_{timestamp}.{export_format}"
